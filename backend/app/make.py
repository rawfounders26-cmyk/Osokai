"""Artifact builders — Osok-AI DELIVERS files, not just words.
Everything lands in workspace/ (visible in Files tab + Command Port)."""
import json, os, subprocess, sys

try:
    WS = os.path.join(os.path.dirname(__file__), "..", "..", "workspace")
except Exception:
    WS = "workspace"
WS = os.path.normpath(WS)
os.makedirs(WS, exist_ok=True)

def _path(name: str) -> str:
    safe = "".join(c for c in name if c.isalnum() or c in ("-", "_", ".", " ")).strip() or "osokai-file"
    return os.path.join(WS, safe)

def make_pptx(title: str, slides_json: str) -> str:
    """slides_json: [{"heading": "...", "bullets": ["..."]}]"""
    from pptx import Presentation
    from pptx.util import Pt
    prs = Presentation()
    t = prs.slides.add_slide(prs.slide_layouts[0])
    t.shapes.title.text = title
    try:
        slides = json.loads(slides_json) if isinstance(slides_json, str) else slides_json
    except Exception:
        slides = []
    for s in slides:
        sl = prs.slides.add_slide(prs.slide_layouts[1])
        sl.shapes.title.text = s.get("heading", "")
        tf = sl.placeholders[1].text_frame
        tf.clear()
        for b in s.get("bullets", [])[:6]:
            p = tf.add_paragraph()
            p.text = str(b)
            p.level = 0
    fp = _path(title[:40] + ".pptx")
    prs.save(fp)
    return fp

def make_xlsx(name: str, sheets_json: str) -> str:
    """sheets_json: {"Sheet1": [["h1","h2"], ["a","b"]]}"""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    try:
        sheets = json.loads(sheets_json) if isinstance(sheets_json, str) else sheets_json
    except Exception:
        sheets = {"Sheet1": []}
    wb = Workbook()
    wb.remove(wb.active)
    for sname, rows in (sheets or {"Sheet1": []}).items():
        ws = wb.create_sheet(str(sname)[:30])
        for r in rows or []:
            ws.append([str(c) for c in r])
        for c in ws[1]:
            c.font = Font(bold=True)
    fp = _path(name if name.endswith(".xlsx") else name + ".xlsx")
    wb.save(fp)
    return fp

def write_file(name: str, content: str) -> str:
    fp = _path(name)
    open(fp, "w", encoding="utf-8").write(content)
    return fp

def make_dir(name: str) -> str:
    """Create a folder in workspace/ (same store the mobile Files screen shows)."""
    safe = "".join(c for c in name if c.isalnum() or c in ("-", "_", " ")).strip() or "osokai-folder"
    fp = os.path.join(WS, safe)
    os.makedirs(fp, exist_ok=True)
    return fp

def run_tests(target: str = "") -> str:
    """Run pytest on workspace code. target: file or '' for all."""
    t = _path(target) if target else WS
    try:
        r = subprocess.run([sys.executable, "-m", "pytest", t, "-q"], capture_output=True, text=True, timeout=120)
        return ((r.stdout or "") + ("\n" + r.stderr if r.stderr else ""))[:3000]
    except subprocess.TimeoutExpired:
        return "(tests timed out)"
    except Exception as e:
        return f"(test error: {e})"

def fetch_url(url: str, max_chars: int = 6000) -> str:
    import re
    import httpx
    try:
        from app.system_tools import clean_url
    except ImportError:
        from system_tools import clean_url
    url = clean_url(url)
    try:
        r = httpx.get(url, headers={"User-Agent": "Mozilla/5.0"}, timeout=25, follow_redirects=True)
        txt = re.sub(r"<script.*?</script>|<style.*?</style>", " ", r.text, flags=re.S | re.I)
        txt = re.sub(r"<[^>]+>", " ", txt)
        return re.sub(r"\s+", " ", txt)[:max_chars]
    except Exception as e:
        return f"(fetch error: {e})"
